#!/usr/bin/env python3
"""Turn the course schedule spreadsheet into a week-by-week timetable.

`Schedule.xlsx` holds one sheet: a column of half-hour slots, five day columns,
and a `WEEK n` row starting each week. An entry that runs for more than half an
hour is either a merged cell or a cell followed by `↳` continuation markers;
both become one block spanning those rows.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

CONTINUATION = "↳"
WEEK = re.compile(r"^\s*WEEK\b", re.I)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _breakable(value: str) -> str:
    """Escape a label and let it wrap after slashes rather than mid-word."""
    return html.escape(value).replace("/", "/<wbr>")


def _rgb(color) -> str:
    """An openpyxl colour as #rrggbb, or "" for theme/indexed/unset colours."""
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return ""
    argb = color.rgb.upper()
    if len(argb) != 8 or argb[:2] == "00":  # fully transparent
        return ""
    return "#" + argb[2:].lower()


def _style(cell) -> str:
    """The cell's fill, text colour and weight, as an inline style attribute."""
    rules = []
    if cell.fill is not None and cell.fill.fill_type == "solid":
        bg = _rgb(cell.fill.fgColor)
        if bg:
            rules.append(f"background:{bg}")
    fg = _rgb(cell.font.color) if cell.font is not None else ""
    if fg:
        rules.append(f"color:{fg}")
    if cell.font is not None and cell.font.b:
        rules.append("font-weight:600")
    return f' style="{";".join(rules)}"' if rules else ""


def _spans(ws) -> dict[tuple[int, int], int]:
    """Top-left cell of each merged block -> how many rows it covers."""
    return {
        (r.min_row, r.min_col): r.max_row - r.min_row + 1
        for r in ws.merged_cells.ranges
    }


def _covered(ws) -> set[tuple[int, int]]:
    """Every cell hidden underneath a merged block."""
    out = set()
    for r in ws.merged_cells.ranges:
        for row in range(r.min_row, r.max_row + 1):
            for col in range(r.min_col, r.max_col + 1):
                if (row, col) != (r.min_row, r.min_col):
                    out.add((row, col))
    return out


def _weeks(ws) -> list[tuple[str, list[str], list[int]]]:
    """Split the sheet into (week label, day headings, row numbers)."""
    weeks: list[tuple[str, list[str], list[int]]] = []
    for row in range(1, ws.max_row + 1):
        label = _text(ws.cell(row, 1).value)
        if not WEEK.match(label):
            continue
        days = [_text(ws.cell(row, col).value) for col in range(2, ws.max_column + 1)]
        body: list[int] = []
        for r in range(row + 1, ws.max_row + 1):
            if WEEK.match(_text(ws.cell(r, 1).value)):
                break
            if _text(ws.cell(r, 1).value):
                body.append(r)
        weeks.append((label, days, body))
    return weeks


def _header_row(ws, label: str) -> int:
    for row in range(1, ws.max_row + 1):
        if _text(ws.cell(row, 1).value) == label:
            return row
    return 1


def _grid(ws, rows: list[int], cols: range) -> dict[tuple[int, int], tuple[str, int]]:
    """Map each event to the cell it starts in and the rows it occupies."""
    spans, covered = _spans(ws), _covered(ws)
    placed: dict[tuple[int, int], tuple[str, int]] = {}
    absorbed: set[tuple[int, int]] = set()

    for col in cols:
        for i, row in enumerate(rows):
            if (row, col) in covered or (row, col) in absorbed:
                continue
            value = _text(ws.cell(row, col).value)
            if not value or value == CONTINUATION:
                continue
            span = spans.get((row, col), 1)
            # A block can be extended by `↳` rows written under it by hand.
            nxt = i + span
            while nxt < len(rows):
                below = rows[nxt]
                if _text(ws.cell(below, col).value) != CONTINUATION:
                    break
                absorbed.add((below, col))
                span += 1
                nxt += 1
            placed[(row, col)] = (value, min(span, len(rows) - i))
    return placed


def _trim(rows: list[int], placed, cols: range) -> list[int]:
    """Drop the empty slots above the first entry and below the last."""
    busy = set()
    for (row, _col), (_value, span) in placed.items():
        start = rows.index(row)
        busy.update(rows[start:start + span])
    if not busy:
        return rows
    return rows[rows.index(min(busy)):rows.index(max(busy)) + 1]


def render(path: Path) -> str:
    """Build the HTML for every week in the workbook."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    cols = range(2, ws.max_column + 1)
    out: list[str] = []

    for n, (label, days, rows) in enumerate(_weeks(ws), start=1):
        placed = _grid(ws, rows, cols)
        rows = _trim(rows, placed, cols)
        wid = f"week-{n}"
        header = _header_row(ws, label)
        corner = _style(ws.cell(header, 1))
        head = "".join(
            f"<th{_style(ws.cell(header, col))}>{html.escape(days[col - 2])}</th>"
            for col in cols
            if days[col - 2]
        )
        body: list[str] = []
        skip: dict[int, int] = {col: 0 for col in cols}

        for row in rows:
            cells: list[str] = []
            for col in cols:
                if not days[col - 2]:
                    continue
                if skip[col]:
                    skip[col] -= 1
                    continue
                entry = placed.get((row, col))
                cell = ws.cell(row, col)
                if entry:
                    value, span = entry
                    skip[col] = span - 1
                    attr = f' rowspan="{span}"' if span > 1 else ""
                    cells.append(
                        f'<td class="slot"{attr}{_style(cell)}>'
                        f"<span>{_breakable(value)}</span></td>"
                    )
                else:
                    # Empty cells keep their fill too: the lunch hour is shaded.
                    cells.append(f'<td class="free"{_style(cell)}></td>')
            time_cell = ws.cell(row, 1)
            time = html.escape(_text(time_cell.value))
            body.append(
                f'<tr><th scope="row"{_style(time_cell)}>{time}</th>{"".join(cells)}</tr>'
            )

        out.append(
            f'<section aria-labelledby="{wid}">\n'
            f'<h2 class="sec" id="{wid}"><span class="num">{n:02d}</span>'
            f"<span>{html.escape(label.title())}</span></h2>\n"
            f'<div class="tablewrap"><table class="timetable">\n'
            f"<thead><tr><th{corner}>Time</th>{head}</tr></thead>\n"
            f'<tbody>{"".join(body)}</tbody>\n'
            "</table></div>\n</section>"
        )

    return "\n\n".join(out)


def nav(path: Path) -> str:
    """The contents rail for the schedule page: one entry per week."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    items = "".join(
        f'<li><a href="#week-{n}"><span class="num">{n}</span>'
        f"<span>{html.escape(label.title())}</span></a></li>"
        for n, (label, _days, _rows) in enumerate(_weeks(ws), start=1)
    )
    return f"<ol>{items}</ol>"
